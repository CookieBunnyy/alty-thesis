from decimal import Decimal

from docx import Document
from io import BytesIO

from app.services.document_extraction import (
    classify_document,
    extract_labeled_fields,
    extract_text,
    normalize_fields,
)


def test_labeled_fields_normalize_common_document_values():
    fields = extract_labeled_fields(
        "Listing ID: PROP-TEST-0001\n"
        "Property Title: Sample Home\n"
        "Price Total: PHP 3,500,000.00\n"
        "Has Garage: Yes\n"
        "Amenities: Pool, Garden"
    )

    normalized = normalize_fields(fields)
    assert normalized["listing_id"] == "PROP-TEST-0001"
    assert normalized["price_total"] == Decimal("3500000.00")
    assert normalized["has_garage"] is True
    assert normalized["amenities"] == ["Pool", "Garden"]


def test_classification_uses_content_not_misleading_filename():
    text = (
        "Listing ID: PROP-TEST-0001\n"
        "Property Title: Sample Home\n"
        "Category: Townhouse\n"
        "Price Total: 3500000"
    )

    assert classify_document(text, "sale_agreement_test.docx") == "PROPERTY_INFORMATION"


def test_transaction_classification_requires_party_and_agent_fields():
    text = (
        "Reservation Agreement\n"
        "Client Full Name: Taylor Example\n"
        "Listing ID: PROP-TEST-0001\n"
        "Agent ID: AGT-0006\n"
        "Reservation Amount: 50000"
    )

    assert classify_document(text) == "RESERVATION_AGREEMENT"


def test_reservation_id_label_is_a_content_classification_signal():
    text = (
        "Client Full Name: Taylor Example\n"
        "Listing ID: PROP-TEST-0001\n"
        "Agent ID: AGT-0006\n"
        "Reservation ID: RES-TEST-0001\n"
        "Transaction Date: 2026-09-12\n"
        "Amount: 50000"
    )

    assert classify_document(text) == "RESERVATION_AGREEMENT"


def test_docx_text_extraction_reads_paragraph_labels():
    document = Document()
    document.add_paragraph("Agent ID: AGT-TEST-0001")
    document.add_paragraph("Full Name: Andrea Reyes")
    content = BytesIO()
    document.save(content)

    text = extract_text("agent.docx", content.getvalue())

    assert "Agent ID: AGT-TEST-0001" in text
    assert "Full Name: Andrea Reyes" in text


def test_unsupported_extension_is_reported():
    try:
        extract_text("notes.txt", b"Listing ID: 1")
    except ValueError as error:
        assert "Unsupported document format" in str(error)
    else:
        raise AssertionError("unsupported formats must not be parsed")